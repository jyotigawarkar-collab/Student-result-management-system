import openpyxl
import os

FILE_NAME = "student_results.xlsx"

# 1. MAIN MENU
def menu():
    while True:
        print("\n========================================")
        print("      STUDENT RESULT MANAGEMENT")
        print("========================================")
        print("1. Add Student Result")
        print("2. Get Student Result")
        print("3. Show All Student Data")
        print("4. Exit")

        choice = input("Enter your choice: ")

        if choice == "1":
            add_student()

        elif choice == "2":
            get_result()

        elif choice == "3":
            show_all_data()

        elif choice == "4":
            print("Thank you for using the system<<3!")
            break

        else:
            print("please enter valid choice !")


# 2. ADD STUDENT
def add_student():
    print("\n----- Add Student Result -----")

    roll_no = input("Enter Roll No: ")
    name = input("Enter Student Name: ")
    course = input("Enter Course/Class: ")

    print("\nEnter marks of 5 subjects:")
    
    marks = []

    for i in range(1, 6):
        mark = int(input(f"Subject {i}: "))
        marks.append(mark)

    total, percentage, grade, status = calculate_result(marks)

    workbook = openpyxl.load_workbook(FILE_NAME)
    sheet = workbook.active

    sheet.append([
        roll_no,
        name,
        course,
        marks[0],
        marks[1],
        marks[2],
        marks[3],
        marks[4],
        total,
        percentage,
        grade,
        status
    ])

    workbook.save(FILE_NAME)

    print("\n----- Result -----")
    print("Roll No:", roll_no)
    print("Name:", name)
    print("Class:", course)
    print("Total:", total)
    print("Percentage:", percentage)
    print("Grade:", grade)
    print("Status:", status)
    print("\nResult saved successfully in Excel!")


# 3. CALCULATE RESULT
def calculate_result(marks):
    total = sum(marks)
    percentage = total / len(marks)

    if percentage >= 80:
        grade = "A"
    elif percentage >= 60:
        grade = "B"
    elif percentage >= 50:
        grade = "C"
    elif percentage >= 40:
        grade = "D"
    else:
        grade = "F"

    if all(mark >= 35 for mark in marks) and percentage >= 40:
        status = "PASS"
    else:
        status = "FAIL"

    return total, percentage, grade, status


# 4. GET STUDENT RESULT
def get_result():
    print("\n----- Get Student Result -----")

    roll_no = input("Enter Roll No: ")

    workbook = openpyxl.load_workbook(FILE_NAME)
    sheet = workbook.active

    found = False

    for row in sheet.iter_rows(min_row=2, values_only=True):

        if str(row[0]) == roll_no:

            print("\n--------------------------------")
            print("        STUDENT RESULT")
            print("--------------------------------")
            print("Roll No    :", row[0])
            print("Name       :", row[1])
            print("Class      :", row[2])
            print("Total      :", row[8])
            print("Percentage :", row[9])
            print("Grade      :", row[10])
            print("Status     :", row[11])
            print("--------------------------------")

            found = True
            break

    if not found:
        print("\nStudent not found!")


# 5. SHOW ALL STUDENT DATA
def show_all_data():
    print("\n----- All Student Data -----")

    workbook = openpyxl.load_workbook(FILE_NAME)
    sheet = workbook.active

    print("\nRoll No | Name | Class | Total | Percentage | Grade | Status")
    print("-" * 65)

    for row in sheet.iter_rows(min_row=2, values_only=True):
        print(
            row[0], "|",
            row[1], "|",
            row[2], "|",
            row[8], "|",
            row[9], "|",
            row[10], "|",
            row[11]
        )


# 6. CREATE EXCEL FILE
def create_excel_file():
    if not os.path.exists(FILE_NAME):
        workbook = openpyxl.Workbook()
        sheet = workbook.active

        sheet.append([
            "Roll No", "Name", "Class",
            "Subject 1", "Subject 2", "Subject 3",
            "Subject 4", "Subject 5",
            "Total", "Percentage", "Grade", "Status"
        ])

        workbook.save(FILE_NAME)


# 7. START PROGRAM
create_excel_file()
menu()
